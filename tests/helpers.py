"""Shared test helpers: sample document builders and a fake AI provider."""

from __future__ import annotations

import base64
import io
import zipfile

import fitz
from docx import Document as DocxDocument

from app.providers import GenerationOutcome

# A 1x1 transparent PNG, enough to make a PDF page "image-only" (scanned).
PNG_1PX = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGA"
    "hKmMIQAAAABJRU5ErkJggg=="
)

DOCX_MIME = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)

SAMPLE_TEXT = (
    "Quarterly Report\n\n"
    "Revenue grew by 12 percent in Q3.\n"
    "- Sales expanded in Europe.\n"
    "- Costs stayed flat.\n"
)


def pdf_bytes(
    heading: str = "Quarterly Report",
    body: str = (
        "Revenue grew by 12 percent in Q3.\n"
        "- Sales expanded in Europe.\n"
        "- Costs stayed flat."
    ),
    *,
    scanned: bool = False,
    encrypted: bool = False,
    blank: bool = False,
) -> bytes:
    doc = fitz.open()
    page = doc.new_page()
    if scanned:
        page.insert_image(fitz.Rect(50, 50, 250, 250), stream=PNG_1PX)
    elif not blank:
        page.insert_text((72, 72), heading, fontsize=22, fontname="helv")
        page.insert_textbox(
            fitz.Rect(72, 100, 540, 320), body, fontsize=11, fontname="helv"
        )
    if encrypted:
        return doc.tobytes(
            encryption=fitz.PDF_ENCRYPT_AES_256, user_pw="secret", owner_pw="secret"
        )
    return doc.tobytes()


def docx_bytes(*, empty: bool = False) -> bytes:
    doc = DocxDocument()
    if not empty:
        doc.add_heading("Introduction", level=1)
        doc.add_paragraph("This report explains the quarterly results.")
        doc.add_paragraph("Sales expanded in Europe.", style="List Bullet")
        doc.add_paragraph("Costs stayed flat.", style="List Bullet")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Metric"
        table.cell(0, 1).text = "Value"
        table.cell(1, 0).text = "Revenue"
        table.cell(1, 1).text = "+12%"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def zip_without_document_xml() -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("hello.txt", "not a docx")
    return buf.getvalue()


class FakeProvider:
    """Duck-typed stand-in for BaseProvider; never touches the network."""

    name = "fake"
    kind = "local"
    model = "fake-1"

    def __init__(self) -> None:
        self.calls: list[dict] = []
        self.error: Exception | None = None
        self.reply = "FAKE EXPLANATION"

    async def chat(self, messages, *, max_tokens: int) -> GenerationOutcome:
        self.calls.append({"messages": messages, "max_tokens": max_tokens})
        if self.error is not None:
            raise self.error
        return GenerationOutcome(
            text=self.reply, provider=self.name, model=self.model, kind=self.kind
        )


def upload(client, filename: str, content: bytes,
           content_type: str = "application/octet-stream"):
    return client.post(
        "/api/documents/extract",
        files={"file": (filename, content, content_type)},
    )


def generate_payload(text: str = "Some document text.", *,
                     format: str = "short_summary", language: str = "en",
                     session_id: str = "test-session-0001") -> dict:
    return {
        "text": text,
        "format": format,
        "language": language,
        "session_id": session_id,
    }
