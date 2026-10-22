"""Document parsing and validation.

Responsibilities:
- enforce upload size limits while streaming (never load unbounded data),
- validate file extension, declared MIME type and real content magic bytes,
- extract text from PDF / DOCX / TXT preserving headings, lists and tables,
- detect scanned (image-only) PDFs and report that OCR is required,
- reject invalid, corrupted and empty documents with stable error codes.

Files are processed fully in memory and never written to disk by the
application, so no uploaded document survives the request.
"""

from __future__ import annotations

import io
import re
import unicodedata
import zipfile
from collections import Counter
from dataclasses import dataclass, field

import fitz  # PyMuPDF
from docx import Document as DocxDocument
from docx.oxml.ns import qn
from docx.opc.exceptions import PackageNotFoundError
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph

from .errors import (
    CorruptedFile,
    EmptyDocument,
    FileTooLarge,
    InvalidFile,
    ScannedDocument,
    UnsupportedFileType,
)

ALLOWED_EXTENSIONS = ("pdf", "docx", "txt")
MAX_FILENAME_LEN = 120

# Declared content types we accept per extension family. Browsers are free to
# send octet-stream for everything, so that value is always allowed; anything
# explicit and foreign (e.g. image/png named .txt) is rejected.
_ACCEPTED_MIME: dict[str, set[str]] = {
    "pdf": {"application/pdf"},
    "docx": {
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/zip",
        "application/x-zip-compressed",
    },
    "txt": {
        "text/plain", "text/markdown", "text/csv",
        "application/msword",  # legacy .doc exported as text; magic bytes decide
    },
}
_GENERIC_MIME = {"", "application/octet-stream", "binary/octet-stream"}

_BULLET_CHARS = "•●▪‣◦·"
_CONTROL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
_PAGE_NUMBER_RE = re.compile(r"^[\s\-–—|·]*\d{1,4}[\s\-–—|·]*$")
_NUMBERED_ITEM_RE = re.compile(r"^\(?\d{1,3}[.)]\s+\S")


@dataclass
class ExtractionResult:
    text: str
    word_count: int = 0
    char_count: int = 0
    headings: list[str] = field(default_factory=list)
    table_count: int = 0
    page_count: int | None = None
    warnings: list[str] = field(default_factory=list)


# --------------------------------------------------------------------------
# Filename / metadata validation
# --------------------------------------------------------------------------
def sanitize_filename(raw: str | None) -> str:
    """Strip any path components and control characters; never trust input."""
    name = (raw or "document").replace("\\", "/").split("/")[-1]
    name = "".join(c for c in name if unicodedata.category(c)[0] != "C")
    name = re.sub(r"\s+", " ", name).strip().lstrip(".").strip()
    if not name:
        return "document"
    if len(name) > MAX_FILENAME_LEN:
        stem, dot, ext = name.rpartition(".")
        if dot and 0 < len(ext) <= 8 and len(stem) > 1:
            name = stem[: MAX_FILENAME_LEN - len(ext) - 1].rstrip(".") + "." + ext
        else:
            name = name[:MAX_FILENAME_LEN]
    return name


def extension_of(filename: str) -> str | None:
    if "." not in filename:
        return None
    return filename.rsplit(".", 1)[1].lower() or None


def validate_file_meta(filename: str, declared_mime: str | None) -> str:
    """Check extension + declared MIME. Returns the normalised extension."""
    ext = extension_of(filename)
    if ext is None or ext not in ALLOWED_EXTENSIONS:
        raise UnsupportedFileType(
            "Unsupported file type. Please upload a PDF, DOCX or TXT file.",
            extra={"allowed_extensions": list(ALLOWED_EXTENSIONS)},
        )
    mime = (declared_mime or "").split(";")[0].strip().lower()
    if mime and mime not in _GENERIC_MIME and mime not in _ACCEPTED_MIME[ext]:
        raise UnsupportedFileType(
            f"The file is declared as '{mime}', which does not match a "
            f".{ext} document. Please upload the original file.",
            extra={"declared_mime": mime, "extension": ext},
        )
    return ext


async def read_upload_limited(upload, max_bytes: int) -> bytes:
    """Stream an UploadFile into memory with a hard size cap."""
    declared = getattr(upload, "size", None)
    if declared is not None and declared > max_bytes:
        raise FileTooLarge(
            f"File is larger than the {_mb(max_bytes)} MB upload limit.",
            extra={"max_bytes": max_bytes, "size_bytes": declared},
        )
    chunks: list[bytes] = []
    total = 0
    while True:
        chunk = await upload.read(1024 * 256)
        if not chunk:
            break
        total += len(chunk)
        if total > max_bytes:
            raise FileTooLarge(
                f"File is larger than the {_mb(max_bytes)} MB upload limit.",
                extra={"max_bytes": max_bytes},
            )
        chunks.append(chunk)
    return b"".join(chunks)


def _mb(n: int) -> float:
    return round(n / (1024 * 1024), 1)


# --------------------------------------------------------------------------
# Text cleanup helpers
# --------------------------------------------------------------------------
def _clean(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x0c", "\n")
    text = _CONTROL_RE.sub("", text)
    return text


def _collapse_ws(text: str) -> str:
    return re.sub(r"[ \t\f\v]+", " ", text).strip()


def _render_table(rows: list[list[str]], table_count_hint: int = 0) -> str:
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    norm = [([_collapse_ws(c.replace("|", "\\|").replace("\n", " / "))
              for c in row] + [""] * (width - len(row)))[:width] for row in rows]
    header, body = norm[0], norm[1:]
    lines = ["| " + " | ".join(header) + " |",
             "| " + " | ".join("---" for _ in header) + " |"]
    lines += ["| " + " | ".join(r) + " |" for r in body]
    return "\n".join(lines)


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------
def _extract_pdf(data: bytes, filename: str) -> ExtractionResult:
    if b"%PDF" not in data[:1024]:
        raise InvalidFile(
            "This file has a .pdf extension but is not a valid PDF. "
            "It may have been renamed or corrupted."
        )
    try:
        doc = fitz.open(stream=data, filetype="pdf")
    except Exception as exc:  # PyMuPDF raises several types for bad data
        raise CorruptedFile("The PDF could not be opened. It may be damaged or incomplete.") from exc

    try:
        if doc.needs_pass:
            raise InvalidFile("This PDF is password-protected. Please upload an unlocked copy.")
        if doc.page_count == 0:
            raise EmptyDocument("This PDF contains no pages.")

        body_size, has_images = _pdf_font_stats(doc)

        parts: list[str] = []
        headings: list[str] = []
        table_count = 0
        total_chars = 0

        for page in doc:
            rendered, page_headings, page_tables, page_chars = _render_pdf_page(
                page, body_size
            )
            headings.extend(page_headings)
            table_count += page_tables
            total_chars += page_chars
            if rendered:
                parts.append(rendered)
            if not has_images:
                has_images = bool(page.get_images(full=True))

        text = _collapse_blank("\n\n".join(parts))

        if total_chars < 40 and has_images:
            raise ScannedDocument(
                "This PDF appears to be scanned (image-only): almost no "
                "extractable text was found. Run OCR on the file first "
                "(e.g. 'ocrmypdf in.pdf out.pdf' or any scanner app with "
                "text recognition), then upload it again."
            )
        if total_chars < 1:
            raise EmptyDocument("No text could be extracted from this PDF.")

        result = _result(text, page_count=doc.page_count)
        result.headings = _unique(headings)[:12]
        result.table_count = table_count
        if table_count:
            result.warnings.append(
                f"{table_count} table(s) were converted to plain-text tables."
            )
        return result
    finally:
        doc.close()


def _pdf_font_stats(doc: fitz.Document) -> tuple[float, bool]:
    """Estimate the dominant body font size from a sample of pages."""
    weights: Counter[float] = Counter()
    has_images = False
    sample = min(doc.page_count, 5)
    for i in range(sample):
        page = doc[i]
        if not has_images:
            has_images = bool(page.get_images(full=True))
        try:
            info = page.get_text("dict")
        except Exception:
            continue
        for block in info.get("blocks", []):
            if block.get("type") != 0:
                continue
            for line in block.get("lines", []):
                for span in line.get("spans", []):
                    size = round(float(span.get("size", 0)) * 2) / 2
                    if size > 0:
                        weights[size] += len(span.get("text", "").strip())
    if not weights:
        return 11.0, has_images
    return max(weights, key=weights.get), has_images  # type: ignore[return-value]


def _render_pdf_page(page, body_size: float) -> tuple[str, list[str], int, int]:
    """One page -> (text, headings, table_count, char_count)."""
    tables: list[tuple[float, str]] = []
    table_boxes: list[tuple[float, float, float, float]] = []
    try:
        finder = page.find_tables()
        found = list(getattr(finder, "tables", []) or [])
        if not found:
            try:
                found = list(finder)
            except TypeError:
                found = []
    except Exception:
        found = []
    for tbl in found:
        try:
            rows = tbl.extract()
        except Exception:
            continue
        if not rows:
            continue
        rendered = _render_table([[c if c else "" for c in r] for r in rows])
        if rendered.replace("|", "").replace("-", "").replace(" ", "").strip():
            tables.append((float(tbl.bbox[1]), rendered))
            table_boxes.append(tuple(float(v) for v in tbl.bbox))  # type: ignore[arg-type]

    try:
        info = page.get_text("dict")
    except Exception:
        return "", [], len(tables), 0

    items: list[tuple[float, str, bool]] = []  # (y, text, is_heading)
    headings: list[str] = []
    chars = 0

    for block in info.get("blocks", []):
        if block.get("type") != 0:
            continue
        if _inside_any(block.get("bbox", (0, 0, 0, 0)), table_boxes):
            continue
        lines_out: list[str] = []
        block_chars = 0
        weighted_size = 0.0
        bold = False
        line_texts: list[str] = []
        for line in block.get("lines", []):
            raw = "".join(span.get("text", "") for span in line.get("spans", []))
            line_texts.append(raw)
            for span in line.get("spans", []):
                t = span.get("text", "")
                weighted_size += float(span.get("size", 0)) * len(t)
                block_chars += len(t)
                if int(span.get("flags", 0)) & 16:
                    bold = True
        if not line_texts or block_chars == 0:
            continue
        avg_size = weighted_size / block_chars if block_chars else body_size

        joined = _collapse_ws(" ".join(line_texts))
        if not joined or _PAGE_NUMBER_RE.match(joined):
            continue

        chars += block_chars
        is_heading = (
            len(line_texts) <= 2
            and len(joined) <= 110
            and (
                avg_size >= body_size * 1.12
                or (bold and avg_size >= body_size * 1.02)
            )
        )
        if is_heading:
            level = 1 if avg_size >= body_size * 1.6 else (
                2 if avg_size >= body_size * 1.35 else 3
            )
            text = "#" * level + " " + joined
            headings.append(joined)
        else:
            text = "\n".join(_render_lines(line_texts))
        items.append((float(block.get("bbox", (0, 0, 0, 0))[1]), text, is_heading))

    for y, rendered in tables:
        items.append((y, rendered, False))

    items.sort(key=lambda it: it[0])
    body = "\n\n".join(text for _, text, _ in items)
    return body, headings, len(tables), chars


def _render_lines(lines: list[str]) -> list[str]:
    out: list[str] = []
    for raw in lines:
        text = _collapse_ws(raw)
        if not text:
            continue
        if text[0] in _BULLET_CHARS:
            text = "- " + text[1:].lstrip(" \t:·")
        out.append(text)
    # merge wrapped lines: a line that does not look like a new item continues
    merged: list[str] = []
    for text in out:
        starts_item = text.startswith("- ") or _NUMBERED_ITEM_RE.match(text)
        if merged and not starts_item and not merged[-1].startswith("- ") \
                and not _NUMBERED_ITEM_RE.match(merged[-1][-80:] if merged else "") \
                and not merged[-1].startswith("#"):
            if len(merged[-1]) < 240 and not merged[-1].endswith((".", "!", "?", ":", ";")):
                merged[-1] = merged[-1] + " " + text
                continue
        merged.append(text)
    return merged


def _inside_any(bbox, boxes) -> bool:
    x0, y0, x1, y1 = bbox
    for tx0, ty0, tx1, ty1 in boxes:
        if x0 >= tx0 - 3 and y0 >= ty0 - 3 and x1 <= tx1 + 3 and y1 <= ty1 + 3:
            return True
    return False


# --------------------------------------------------------------------------
# DOCX
# --------------------------------------------------------------------------
def _extract_docx(data: bytes, filename: str) -> ExtractionResult:
    if not data.startswith(b"PK"):
        raise InvalidFile(
            "This file has a .docx extension but is not a valid DOCX file."
        )
    try:
        if not zipfile.is_zipfile(io.BytesIO(data)):
            raise InvalidFile(
                "This file has a .docx extension but is not a valid DOCX file."
            )
        doc = DocxDocument(io.BytesIO(data))
    except (InvalidFile, CorruptedFile, EmptyDocument, UnsupportedFileType):
        raise
    except (zipfile.BadZipFile, PackageNotFoundError, KeyError, ValueError) as exc:
        raise CorruptedFile(
            "The DOCX file is corrupted or incomplete and could not be read."
        ) from exc
    except Exception as exc:
        raise CorruptedFile("The DOCX file could not be parsed.") from exc

    blocks: list[str] = []
    headings: list[str] = []
    pending_list: list[str] = []
    table_count = 0

    def flush_list() -> None:
        nonlocal pending_list
        if pending_list:
            blocks.append("\n".join(pending_list))
            pending_list = []

    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            para = DocxParagraph(child, doc)
            text = _collapse_ws(_clean(para.text or ""))
            if not text:
                flush_list()
                continue
            style = (para.style.name if para.style is not None else "") or ""
            style_l = style.lower()
            if style_l.startswith("heading") or style_l == "title" or style_l.startswith("title"):
                flush_list()
                m = re.search(r"(\d)", style)
                level = min(int(m.group(1)) if m else 2, 4)
                blocks.append("#" * max(level, 1) + " " + text)
                headings.append(text)
            elif _is_docx_list(child, style_l):
                pending_list.append("- " + text)
            else:
                flush_list()
                blocks.append(text)
        elif child.tag == qn("w:tbl"):
            flush_list()
            try:
                table = DocxTable(child, doc)
                rows = [[_clean(cell.text or "") for cell in row.cells]
                        for row in table.rows]
            except Exception:
                continue
            rendered = _render_table(rows)
            if rendered:
                blocks.append(rendered)
                table_count += 1
    flush_list()

    text = _collapse_blank("\n\n".join(blocks))
    if not text.strip():
        raise EmptyDocument("This document contains no extractable text.")

    result = _result(text)
    result.headings = _unique(headings)[:12]
    result.table_count = table_count
    if table_count:
        result.warnings.append(
            f"{table_count} table(s) were converted to plain-text tables."
        )
    return result


def _is_docx_list(paragraph_element, style_lower: str = "") -> bool:
    """Lists are either directly numbered or styled with a List style."""
    if style_lower.startswith("list"):
        return True
    ppr = paragraph_element.find(qn("w:pPr"))
    if ppr is None:
        return False
    return ppr.find(qn("w:numPr")) is not None


# --------------------------------------------------------------------------
# TXT
# --------------------------------------------------------------------------
def _extract_txt(data: bytes, filename: str) -> ExtractionResult:
    is_utf16 = data[:2] in (b"\xff\xfe", b"\xfe\xff")
    if not is_utf16 and b"\x00" in data[:8000]:
        raise InvalidFile("This .txt file contains binary data and cannot be read as text.")

    text: str | None = None
    if is_utf16:
        try:
            text = data.decode("utf-16")
        except UnicodeDecodeError:
            text = None
    if text is None:
        for enc in ("utf-8-sig", "utf-8", "cp1252"):
            try:
                text = data.decode(enc)
                break
            except UnicodeDecodeError:
                continue
    if text is None:
        text = data.decode("latin-1")
        printable = sum(1 for c in text[:5000] if c.isprintable() or c in "\n\r\t")
        if printable / max(1, min(len(text), 5000)) < 0.85:
            raise InvalidFile("This .txt file does not contain readable text.")

    text = _collapse_blank(_clean(text))
    if not text.strip():
        raise EmptyDocument("This file is empty or contains no readable text.")

    result = _result(text)
    headings = [ln.lstrip("#").strip() for ln in text.splitlines()
                if ln.startswith("# ") and len(ln) < 100]
    result.headings = _unique(headings)[:12]
    return result


# --------------------------------------------------------------------------
# Dispatcher
# --------------------------------------------------------------------------
def extract_text(filename: str, data: bytes) -> ExtractionResult:
    """Validate content and extract text for the given extension."""
    if not data:
        raise EmptyDocument("The uploaded file is empty.")
    ext = extension_of(filename)
    if ext == "pdf":
        return _extract_pdf(data, filename)
    if ext == "docx":
        return _extract_docx(data, filename)
    if ext == "txt":
        return _extract_txt(data, filename)
    raise UnsupportedFileType("Unsupported file type. Please upload a PDF, DOCX or TXT file.")


def extract_upload(filename: str, declared_mime: str | None, data: bytes) -> ExtractionResult:
    """Full pipeline: metadata validation -> magic validation -> extraction."""
    name = sanitize_filename(filename)
    validate_file_meta(name, declared_mime)
    return extract_text(name, data)


# --------------------------------------------------------------------------
# Helpers
# --------------------------------------------------------------------------
def _collapse_blank(text: str) -> str:
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _result(text: str, page_count: int | None = None) -> ExtractionResult:
    text = text.strip()
    return ExtractionResult(
        text=text,
        word_count=len(text.split()),
        char_count=len(text),
        page_count=page_count,
    )


def _unique(items: list[str]) -> list[str]:
    seen: set[str] = set()
    out: list[str] = []
    for item in items:
        key = item.lower()
        if key not in seen:
            seen.add(key)
            out.append(item)
    return out
